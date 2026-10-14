"""
CODS Excel Service - Robust Court Schedule Spreadsheet Ingestion
Provides column heuristic mapping, date normalization, Buddhist era year conversions,
and table extraction for judicial proceedings.
100% Offline (Air-Gapped) - Zero external network interactions.
"""

import re
import os
from typing import Dict, Any, List, Tuple, Optional
import openpyxl

COMMON_HEADERS: Dict[str, List[str]] = {
    'case_black_no': [
        'หมายเลขคดีดำ', 'คดีหมายเลขดำ', 'เลขคดีดำ', 'คดีดำ', 'เลขดำ', 'หมายเลขดำ',
        'คดีดำที่', 'หมายเลขคดีดำที่', 'black_no', 'case_black', 'black'
    ],
    'case_red_no': [
        'หมายเลขคดีแดง', 'คดีหมายเลขแดง', 'เลขคดีแดง', 'คดีแดง', 'เลขแดง', 'หมายเลขแดง',
        'คดีแดงที่', 'หมายเลขคดีแดงที่', 'red_no', 'case_red', 'red'
    ],
    'hearing_date': [
        'วันนัดพิจารณา', 'วันที่นัดพิจารณา', 'วันที่นัด', 'วันนัด', 'วันที่', 'วันเดือนปี',
        'วันนัดพิจารณาคดี', 'วันเวลาที่นัด', 'วันเวลา', 'วัน/เวลา'
    ],
    'hearing_day': [
        'วัน'
    ],
    'hearing_month': [
        'เดือน'
    ],
    'hearing_year': [
        'ปี', 'พ.ศ.', 'ปี พ.ศ.', 'ปีพ.ศ.', 'พ.ศ'
    ],
    'hearing_time': [
        'เวลานัด', 'เวลาเริ่ม', 'เวลาที่นัด', 'เวลา'
    ],
    'plaintiff_name': [
        'ชื่อโจทก์', 'โจทก์', 'ผู้ร้อง', 'โจทก์/ผู้ร้อง', 'ผู้กล่าวหา'
    ],
    'defendant_name': [
        'ชื่อจำเลย', 'จำเลย', 'ผู้คัดค้าน', 'จำเลย/ผู้คัดค้าน', 'ผู้ถูกกล่าวหา'
    ],
    'case_type': [
        'ประเภทคดี', 'ประเภทความ', 'ประเภทคดีความ', 'ลักษณะคดี', 'ฐานความผิด'
    ],
    'case_subject': [
        'เรื่อง', 'ข้อหา', 'ฐานความผิด', 'เรื่องที่ฟ้อง', 'มูลคดี'
    ],
    'court_name': [
        'ชื่อศาล', 'ศาล', 'ศาลที่นัด'
    ],
    'hearing_purpose': [
        'นัดมาทำไม', 'เหตุที่นัด', 'สาเหตุที่นัด', 'นัดเรื่อง', 'นัดเพื่อ', 'การนัด', 'นัดมา', 'ข้อหานัด'
    ],
    'appointment_date': [
        'นัดมาทำไม', 'เหตุที่นัด', 'สาเหตุที่นัด', 'นัดเรื่อง', 'นัดเพื่อ', 'การนัด', 'นัดมา', 'ข้อหานัด'
    ],
    'judge_1_name': [
        'เจ้าของสำนวน', 'ผู้พิพากษาเจ้าของสำนวน', 'ผู้พิพากษาคนที่ 1', 'ผู้พิพากษา 1'
    ],
    'judge_2_name': [
        'องค์คณะ', 'ผู้พิพากษาองค์คณะ', 'ผู้พิพากษาคนที่ 2', 'ผู้พิพากษา 2'
    ],
    'judge_name': [
        'ผู้พิพากษา', 'บัลลังก์', 'ห้องพิจารณา', 'ห้องพิจารณาคดี'
    ],
}

SEQUENCE_NAMES = {'ที่', 'ลำดับ', 'ลำดับที่', 'ข้อ', 'no', 'no.', 'item', 'seq', 'index', 'number'}

HEADER_KEYWORDS = [
    'ดำ', 'แดง', 'คดี', 'โจทก์', 'จำเลย', 'นัด', 'เวลา', 'ลำดับ', 'ที่',
    'ผู้ร้อง', 'ผู้คัดค้าน', 'บัลลังก์', 'ผู้พิพากษา', 'ประเภท', 'เรื่อง'
]

THAI_MONTHS = [
    'มกราคม', 'กุมภาพันธ์', 'มีนาคม', 'เมษายน', 'พฤษภาคม', 'มิถุนายน',
    'กรกฎาคม', 'สิงหาคม', 'กันยายน', 'ตุลาคม', 'พฤศจิกายน', 'ธันวาคม'
]

THAI_SHORT_MONTHS = {
    'ม.ค.': 'มกราคม', 'ก.พ.': 'กุมภาพันธ์', 'มี.ค.': 'มีนาคม', 'เม.ย.': 'เมษายน',
    'พ.ค.': 'พฤษภาคม', 'มิ.ย.': 'มิถุนายน', 'ก.ค.': 'กรกฎาคม', 'ส.ค.': 'สิงหาคม',
    'ก.ย.': 'กันยายน', 'ต.ค.': 'ตุลาคม', 'พ.ย.': 'พฤศจิกายน', 'ธ.ค.': 'ธันวาคม'
}


def guess_column_mapping(headers: List[str], rows_data: Optional[List[Dict[str, Any]]] = None) -> Dict[str, str]:
    """
    Intelligently map court schedule spreadsheet columns to standardized case fields
    using a multi-pass heuristic matching algorithm.
    """
    mapping: Dict[str, str] = {k: '' for k in COMMON_HEADERS.keys()}
    used_headers = set()

    # Pass 1: Exact matches (Priority 1 to avoid false substring matches)
    for field, candidates in COMMON_HEADERS.items():
        for candidate in candidates:
            cand_norm = candidate.replace(' ', '').lower()
            for header in headers:
                if header in used_headers:
                    continue
                header_norm = header.replace(' ', '').lower()
                if header_norm == cand_norm:
                    mapping[field] = header
                    used_headers.add(header)
                    break
            if mapping[field]:
                break

    # Pass 2: High-Confidence Court Prefix/Suffix Matches (Safe Substring)
    if not mapping['case_black_no']:
        for header in headers:
            if header in used_headers:
                continue
            h_norm = header.replace(' ', '').lower()
            # Match 'ดำ' but explicitly NOT 'ดำเนิน' (e.g. การดำเนินคดี)
            if re.search(r'ดำ(?!เนิน)', h_norm) and not any(kw in h_norm for kw in ['แดง', 'ดำเนิน']):
                mapping['case_black_no'] = header
                used_headers.add(header)
                break

    if not mapping['case_red_no']:
        for header in headers:
            if header in used_headers:
                continue
            h_norm = header.replace(' ', '').lower()
            if 'แดง' in h_norm and not any(kw in h_norm for kw in ['ดำ', 'ดำเนิน']):
                mapping['case_red_no'] = header
                used_headers.add(header)
                break
            elif re.search(r'\bred\b', h_norm) and 'black' not in h_norm:
                mapping['case_red_no'] = header
                used_headers.add(header)
                break

    # Pass 3: Fuzzy matches on remaining headers
    for field, candidates in COMMON_HEADERS.items():
        if mapping[field]:
            continue
        for candidate in candidates:
            cand_norm = candidate.replace(' ', '').lower()
            for header in headers:
                if header in used_headers:
                    continue
                header_norm = header.replace(' ', '').lower()
                if cand_norm in header_norm or header_norm in cand_norm:
                    if field == 'case_type' and 'คู่ความ' in header_norm:
                        continue  # Avoid matching litigants as case type
                    if field == 'case_black_no' and 'ดำเนิน' in header_norm:
                        continue
                    if field in ('hearing_purpose', 'appointment_date') and header_norm == 'เรื่อง':
                        continue  # 'เรื่อง' is the subject matter/cause of action, NOT the hearing purpose ('นัดมาทำไม')
                    if field == 'court_name':
                        # Exclude attendance, court orders, summons, staff, courtroom, notes
                        if any(kw in header_norm for kw in ['มาศาล', 'คำสั่ง', 'หมาย', 'เจ้าหน้าที่', 'ห้อง', 'บัลลังก์', 'ผล', 'สถานะ']):
                            continue
                        if cand_norm == 'ศาล' and not header_norm.startswith('ศาล'):
                            continue
                    mapping[field] = header
                    used_headers.add(header)
                    break
            if mapping[field]:
                break

    # Pass 4: Content-based inspection on early rows
    if rows_data and len(rows_data) > 0:
        sample_rows = rows_data[:10]

        # Case Black number: look for exactly 1 slash (not dates with 2 slashes)
        if not mapping['case_black_no']:
            for h in headers:
                if h in used_headers:
                    continue
                slash_1_count = 0
                slash_2_count = 0
                for r in sample_rows:
                    val_str = str(r.get(h, '')).strip()
                    if val_str.count('/') == 1 and not re.search(r'\d{1,2}/\d{1,2}/\d{2,4}', val_str):
                        slash_1_count += 1
                    elif val_str.count('/') >= 2:
                        slash_2_count += 1
                if slash_1_count >= 2 and slash_2_count == 0:
                    mapping['case_black_no'] = h
                    used_headers.add(h)
                    break

        # Case Red number: look for red case format
        if not mapping['case_red_no']:
            for h in headers:
                if h in used_headers:
                    continue
                red_match_count = 0
                for r in sample_rows:
                    val_str = str(r.get(h, '')).strip()
                    if 'แดง' in val_str and ('/' in val_str or any(c.isdigit() for c in val_str)):
                        red_match_count += 1
                if red_match_count >= 2:
                    mapping['case_red_no'] = h
                    used_headers.add(h)
                    break

        # Hearing Time: look for colon or point time formats
        if not mapping['hearing_time']:
            for h in headers:
                if h in used_headers:
                    continue
                time_count = sum(
                    1 for r in sample_rows
                    if re.match(r'^\d{1,2}[:.]\d{2}', str(r.get(h, '')).strip())
                )
                if time_count >= 2:
                    mapping['hearing_time'] = h
                    used_headers.add(h)
                    break

        # Court Name: look for columns where multiple rows contain actual court names starting with 'ศาล'
        if not mapping['court_name']:
            for h in headers:
                if h in used_headers:
                    continue
                court_count = 0
                for r in sample_rows:
                    val_str = str(r.get(h, '')).strip()
                    if val_str.startswith('ศาล') and not any(kw in val_str for kw in ['มาศาล', 'คำสั่ง', 'หมาย', 'เจ้าหน้าที่', 'ห้อง']):
                        court_count += 1
                if court_count >= 2:
                    mapping['court_name'] = h
                    used_headers.add(h)
                    break

        # Sync hearing_purpose and appointment_date aliases (always mapped to 'นัดมาทำไม')
        if mapping.get('hearing_purpose') and not mapping.get('appointment_date'):
            mapping['appointment_date'] = mapping['hearing_purpose']
        elif mapping.get('appointment_date') and not mapping.get('hearing_purpose'):
            mapping['hearing_purpose'] = mapping['appointment_date']

    return mapping


def read_excel_data(file_path: str) -> Tuple[List[str], List[Dict[str, Any]]]:
    """
    Read court schedule spreadsheet (.xlsx or .xls) safely, filter empty rows/columns,
    detect header row, and return cleaned headers and row dictionaries.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    raw_data: List[List[Any]] = []

    if ext == '.xlsx':
        wb = openpyxl.load_workbook(file_path, data_only=True)
        try:
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                raw_data.append(list(row))
        finally:
            wb.close()
    elif ext == '.xls':
        import xlrd
        wb = xlrd.open_workbook(file_path)
        try:
            ws = wb.sheet_by_index(0)
            for r in range(ws.nrows):
                row_vals = []
                for c in range(ws.ncols):
                    cell = ws.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            dt = xlrd.xldate.xldate_as_datetime(cell.value, wb.datemode)
                            row_vals.append(dt.strftime("%Y-%m-%d"))
                        except Exception:
                            row_vals.append(str(cell.value))
                    else:
                        row_vals.append(cell.value)
                raw_data.append(row_vals)
        finally:
            wb.release_resources()
    else:
        raise ValueError(f"Unsupported file format: {ext}. Only .xlsx and .xls are supported.")

    if not raw_data:
        return [], []

    # Filter out empty rows
    cleaned_rows: List[List[Any]] = []
    for row in raw_data:
        if any(cell is not None and str(cell).strip() != '' for cell in row):
            cleaned_rows.append(row)

    if not cleaned_rows:
        return [], []

    # Detect header row
    header_idx = -1
    for idx, row in enumerate(cleaned_rows[:15]):
        str_cells = [str(c).strip() for c in row if c is not None and str(c).strip() != '']
        if not str_cells:
            continue
        keyword_matches = sum(1 for cell in str_cells if any(kw in cell for kw in HEADER_KEYWORDS))
        non_empty_count = len(str_cells)
        if (non_empty_count >= 3 and keyword_matches >= 2) or (keyword_matches >= 3):
            header_idx = idx
            break

    if header_idx == -1:
        header_idx = 0

    header_row = cleaned_rows[header_idx]
    data_rows = cleaned_rows[header_idx + 1:]

    # Remove columns that are completely empty in both header and all rows
    non_empty_col_indices = []
    max_len = max(len(r) for r in cleaned_rows)
    for col_idx in range(max_len):
        h_val = header_row[col_idx] if col_idx < len(header_row) else None
        h_has_text = h_val is not None and str(h_val).strip() != ''
        d_has_text = False
        for r in data_rows:
            if col_idx < len(r) and r[col_idx] is not None and str(r[col_idx]).strip() != '':
                d_has_text = True
                break
        if h_has_text or d_has_text:
            non_empty_col_indices.append(col_idx)

    # Build unique clean header names
    headers: List[str] = []
    used_names: Dict[str, int] = {}
    for col_idx in non_empty_col_indices:
        val = header_row[col_idx] if col_idx < len(header_row) else None
        name = str(val).strip() if val is not None and str(val).strip() != '' else f"Column_{col_idx + 1}"
        name = re.sub(r'[\r\n\t]+', ' ', name).strip()
        if name in used_names:
            used_names[name] += 1
            name = f"{name}_{used_names[name]}"
        else:
            used_names[name] = 1
        headers.append(name)

    # Build row dictionaries
    result_rows: List[Dict[str, Any]] = []
    for r in data_rows:
        row_dict: Dict[str, Any] = {}
        has_content = False
        for i, col_idx in enumerate(non_empty_col_indices):
            val = r[col_idx] if col_idx < len(r) else ''
            if val is None:
                val = ''
            elif isinstance(val, float) and val.is_integer():
                val = int(val)
            val_str = str(val).strip()
            row_dict[headers[i]] = val_str
            if val_str:
                has_content = True
        if has_content:
            result_rows.append(row_dict)

    return headers, result_rows


def to_full_buddhist_year(val: Any) -> str:
    """
    Convert 2-digit or 4-digit Gregorian or Buddhist years to standard 4-digit Buddhist Era (B.E.).
    Examples:
      '2569' -> '2569'
      '69'   -> '2569'
      '2026' -> '2569'
      '26'   -> '2569'
    """
    if val is None:
        return ""
    s = str(val).strip()
    digits = re.findall(r'\d+', s)
    if not digits:
        return s

    num = int(digits[0])
    if num >= 2400:
        return str(num)
    elif 1900 <= num <= 2200:
        return str(num + 543)
    elif 20 <= num < 50:
        # 2-digit Gregorian year in current era (e.g. 26 -> 2026 -> 2569)
        return str(2000 + num + 543)
    elif 50 <= num < 100:
        # 2-digit Buddhist Era year (e.g. 69 -> 2569)
        return str(2500 + num)
    elif num < 20:
        return str(2000 + num + 543)

    return s


def format_thai_date(date_val: Any) -> Tuple[str, str, str]:
    """
    Parse date values (datetime, string, ISO format) into Thai (day, month_name, full_be_year).
    Returns (day, month, year) tuple.
    """
    if date_val is None:
        return '', '', ''

    if hasattr(date_val, 'strftime'):
        d = str(date_val.day)
        m = THAI_MONTHS[date_val.month - 1]
        y = str(date_val.year + 543 if date_val.year < 2400 else date_val.year)
        return d, m, y

    s = str(date_val).strip()
    if not s or s in ['-', 'N/A', 'None']:
        return '', '', ''

    # Replace short Thai month names with full names
    for short_m, full_m in THAI_SHORT_MONTHS.items():
        if short_m in s:
            s = s.replace(short_m, full_m)
            break

    # Check if a full Thai month name is present in string
    for full_m in THAI_MONTHS:
        if full_m in s:
            s_clean = re.sub(r'\d{1,2}:\d{1,2}(:\d{1,2})?', '', s).strip()
            parts = [p for p in re.split(r'[-/.\s]+', s_clean) if p]
            d, y = '', ''
            m = full_m
            for p in parts:
                if p == full_m:
                    continue
                if p.isdigit():
                    num = int(p)
                    if num > 1000:
                        y = to_full_buddhist_year(num)
                    elif num > 31 and num < 100:
                        # 2-digit year (e.g. 69)
                        y = to_full_buddhist_year(num)
                    elif not d:
                        d = str(num)
                    elif not y:
                        y = to_full_buddhist_year(num)
            return d, m, to_full_buddhist_year(y) if y else ''

    # Numeric date parsing: remove time portion first (e.g. '00:00:00')
    s_clean = re.sub(r'\d{1,2}:\d{1,2}(:\d{1,2})?', '', s).strip()
    parts = [p for p in re.split(r'[-/.\s]+', s_clean) if p]

    if len(parts) >= 3:
        p0, p1, p2 = parts[0], parts[1], parts[2]
        if p0.isdigit() and p1.isdigit() and p2.isdigit():
            n0, n1, n2 = int(p0), int(p1), int(p2)
            # Format 1: YYYY-MM-DD (e.g. 2026-10-05)
            if n0 > 1000:
                y = to_full_buddhist_year(n0)
                m = THAI_MONTHS[n1 - 1] if 1 <= n1 <= 12 else str(n1)
                d = str(n2)
                return d, m, y
            # Format 2: DD-MM-YYYY (e.g. 05-10-2026 or 05-10-2569 or 05-10-69)
            else:
                d = str(n0)
                m = THAI_MONTHS[n1 - 1] if 1 <= n1 <= 12 else str(n1)
                y = to_full_buddhist_year(n2)
                return d, m, y

    return s, '', ''
